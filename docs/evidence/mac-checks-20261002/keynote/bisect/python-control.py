# python-pptx control deck for the Keynote chart bisect (opf-pptx#162): one slide, one clustered column chart.
# Usage: <venv>/bin/python bisect/python-control.py decks/b-90-python-pptx-column.pptx   (needs python-pptx, XlsxWriter)
import sys
import pptx
import xlsxwriter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
title = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.8)).text_frame
title.text = 'Column (python-pptx)'
title.paragraphs[0].runs[0].font.size = Pt(28)
data = CategoryChartData()
data.categories = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun']
data.add_series('North', [10 + i * 3 for i in range(6)])
data.add_series('South', [14 + ((i * 5) % 9) for i in range(6)])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4), Inches(12), Inches(5.4), data)
prs.save(sys.argv[1])
print(f'python-pptx {pptx.__version__} + XlsxWriter {xlsxwriter.__version__}')
